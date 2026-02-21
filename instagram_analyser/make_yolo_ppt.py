from pptx import Presentation
from pptx.util import Inches

prs = Presentation()

# === 投影片1 標題 ===
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "YOLOv8・YOLOv11 比較實驗報告"
slide.placeholders[1].text = "東海大學 資訊工程學系\n製作者：Zainab（2025年10月）"

# === 投影片2 研究目的 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "1️⃣ 研究目的"
slide.placeholders[1].text = (
    "本研究旨在利用人工智慧自動偵測醫師於X光影像上標註的11個骨骼關鍵點，"
    "進而自動計算用於扁平足判定的角度。\n\n"
    "本週的研究重點如下：\n"
    "・比較 YOLOv8s 與 YOLOv11s 的效能\n"
    "・基於結果判斷 YOLOv8 較適合醫學影像，進一步比較 YOLOv8s 與 YOLOv8l"
)

# === 投影片3 實驗環境 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "2️⃣ 實驗環境"
slide.placeholders[1].text = (
    "・環境：東海大學 SSH 伺服器\n"
    "・GPU：NVIDIA RTX 5090（32GB）\n"
    "・作業系統：Ubuntu 22.04\n"
    "・框架：PyTorch + Ultralytics YOLOv8 / YOLOv11\n"
    "・Python：3.12（虛擬環境）"
)

# === 投影片4 資料集概要 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "3️⃣ 資料集概要"
slide.placeholders[1].text = (
    "・醫師手動於LAT X光影像上標記11個關鍵點\n"
    "・資料集來源：LAT 0-70_100、LAT 71-107_51、LAT 109-139_49、LAT_40\n"
    "・總計：240 張影像\n"
    "・標註工具：VIA（VGG Image Annotator）"
)

# === 投影片5 訓練時間比較 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "4️⃣ 訓練時間比較"
slide.placeholders[1].text = (
    "YOLOv8s：80epoch → 0.029小時（約1.7分鐘）、170epoch → 0.095小時（約5.7分鐘）\n"
    "YOLOv11s：80epoch → 0.034小時、170epoch → 0.107小時\n"
    "YOLOv8l：100epoch → 0.262小時、200epoch → 0.878小時（約52.7分鐘）\n\n"
    "⚙️ YOLOv8s 與 YOLOv11s 屬於高速訓練模型，而 YOLOv8l 雖然耗時較久但精度更高。"
)

# === 投影片6 YOLOv8s vs YOLOv11s ===
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "5️⃣ YOLOv8s 與 YOLOv11s 比較結果"
slide.shapes.add_picture("v8s_results.png", Inches(0.5), Inches(1.5), height=Inches(3.0))
slide.shapes.add_picture("v11s_results.png", Inches(5.0), Inches(1.5), height=Inches(3.0))
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(4.8), Inches(9), Inches(1.5))
txBox.text = (
    "・mAP50-95(P)：v8s=0.982、v11s=0.961\n"
    "・平均角度MAE：v8s=2.68°、v11s=3.96°（v8s 誤差低32%）\n"
    "・平均關鍵點誤差：v8s=2.0px、v11s=3.1px\n"
    "✅ YOLOv8s 在醫學影像中的微小結構偵測表現更佳"
)

# === 投影片7 YOLOv8s vs YOLOv8l ===
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "6️⃣ YOLOv8s 與 YOLOv8l 比較結果"
slide.shapes.add_picture("v8s_results.png", Inches(0.5), Inches(1.5), height=Inches(3.0))
slide.shapes.add_picture("v8l_results.png", Inches(5.0), Inches(1.5), height=Inches(3.0))
txBox = slide.shapes.add_textbox(Inches(0.5), Inches(4.8), Inches(9), Inches(1.5))
txBox.text = (
    "・mAP50-95(P)：v8s=0.982、v8l=0.993\n"
    "・平均角度MAE：v8s=2.68°、v8l=2.39°（v8l 準確度高約11%）\n"
    "・訓練時間：v8l 約為 v8s 的9倍\n"
    "✅ YOLOv8s 適合實務應用，v8l 適合研究與精密分析"
)

# === 投影片8 結果分析 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "7️⃣ 結果分析"
slide.placeholders[1].text = (
    "・YOLOv8 能有效處理邊緣模糊的X光影像，適合醫用關鍵點偵測。\n"
    "・YOLOv11 的輕量化Transformer結構導致微小特徵損失。\n"
    "・未來將考慮導入 ROI 權重損失函數以提升 Navicular 附近精度。"
)

# === 投影片9 未來展望 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "8️⃣ 未來展望"
slide.placeholders[1].text = (
    "・計畫新增100張X光影像進行再訓練。\n"
    "・評估 YOLOv11l 或 YOLOv8/11 混合模型結構。\n"
    "・研究影像前處理技術，如雜訊抑制與對比增強。"
)

# === 投影片10 參考文獻 ===
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "📚 參考文獻"
slide.placeholders[1].text = (
    "1. Ultralytics YOLOv8 / YOLOv11 官方文件\n"
    "2. Bochkovskiy, A. et al. YOLOv4: Optimal Speed and Accuracy of Object Detection (2020)\n"
    "3. Redmon, J. et al. You Only Look Once (CVPR 2016)\n"
    "4. Wang, C.-Y. et al. YOLOv7: Trainable Bag-of-Freebies (arXiv 2022)\n"
    "5. 醫師提供之X光資料與自主VIA標註"
)

prs.save("YOLOv8_v11_比較報告_繁體中文版.pptx")
print("✅ 已成功建立 YOLOv8_v11_比較報告_繁體中文版.pptx")
